#!/usr/bin/env python3
"""Rebuild offline validation and review artifacts for this opportunity intake."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
CSV_SOURCE = ROOT / "oppscout_seed_opportunities.csv"
XLSX_SOURCE = ROOT / "oppscout_seed_opportunities.xlsx"
README = ROOT / "README.md"
XLSX_REVIEWED = ROOT / "oppscout_seed_opportunities_reviewed.xlsx"
CSV_REVIEWED = ROOT / "oppscout_seed_opportunities_reviewed.csv"
REPORT = ROOT / "validation_report.reviewed.json"
HANDOFF = ROOT / "handoff.md"
CHECKED_AT = datetime(2026, 10, 12, 9, 15, 53, tzinfo=timezone.utc)
JSON_COLUMNS = {
    "eligibility_json",
    "required_skills_json",
    "preferred_skills_json",
    "eligible_countries_json",
    "additional_source_urls_json",
    "evidence_json",
    "unresolved_requirements_json",
}
AGENT_COLUMNS = [
    "agent_verdict",
    "agent_checked_at_utc",
    "official_listing_url",
    "value_score_0_100",
    "effort_score_0_100",
    "risk_score_0_100",
    "evidence_files",
]
ENUMS = {
    "action": {"create", "update_existing", "exclude"},
    "category": {"internship", "job", "training", "consultancy", "scholarship", "grant", "fellowship"},
    "applicant_type": {"individual", "organization", "both", "mixed", "unknown"},
    "verification_status": {"pending", "flagged", "verified"},
    "source": {"scraped", "org_submitted", "seed"},
    "deadline_precision": {"date", "datetime", "none", "unknown"},
    "deadline_type": {"fixed", "rolling", "unknown"},
    "region_tier": {"uganda", "east_africa", "africa", "global"},
    "agent_verdict": {"verified", "unresolved", "rejected", "expired"},
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def parse_time(value: str) -> datetime | None:
    if not value:
        return None
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    return parsed.astimezone(timezone.utc) if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


def parse_date(value: str) -> date | None:
    return date.fromisoformat(value) if value else None


def verdict_for(row: dict[str, str]) -> str:
    explicit = parse_time(row["deadline_utc"])
    closing_day = parse_date(row["deadline_date"])
    if row["action"] == "exclude":
        if (closing_day and closing_day < CHECKED_AT.date()) or (explicit and explicit <= CHECKED_AT):
            return "expired"
        notes = f"{row['review_notes']} {row['unresolved_requirements_json']}".lower()
        return "expired" if re.search(r"closed|stale|closing date|deadline .* passed|applications closed", notes) else "rejected"
    if (closing_day and closing_day < CHECKED_AT.date()) or (explicit and explicit <= CHECKED_AT):
        return "expired"
    # No source was rechecked on its provider-owned page in this offline audit.
    return "unresolved"


def main() -> None:
    with CSV_SOURCE.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream, strict=True)
        headers = reader.fieldnames or []
        source_rows = list(reader)
    if len(headers) != 43 or len(source_rows) != 160:
        raise ValueError(f"Unexpected interchange shape: {len(headers)} columns, {len(source_rows)} rows")

    errors: list[str] = []
    parsed_json_cells = 0
    for index, row in enumerate(source_rows, start=2):
        for column in JSON_COLUMNS:
            raw = row.get(column, "")
            if raw:
                try:
                    json.loads(raw)
                    parsed_json_cells += 1
                except json.JSONDecodeError as error:
                    errors.append(f"CSV row {index}: invalid {column}: {error}")
        for column, allowed in ENUMS.items():
            if column in row and row[column] and row[column] not in allowed:
                errors.append(f"CSV row {index}: invalid {column}={row[column]!r}")
        for column in ("source_url", "application_url", "organization_official_url"):
            value = row.get(column, "")
            if value and urlparse(value).scheme != "https":
                errors.append(f"CSV row {index}: non-HTTPS {column}")
        if row["deadline_utc"]:
            try:
                parse_time(row["deadline_utc"])
            except ValueError:
                errors.append(f"CSV row {index}: invalid deadline_utc")
        if row["deadline_date"]:
            try:
                parse_date(row["deadline_date"])
            except ValueError:
                errors.append(f"CSV row {index}: invalid deadline_date")

    if len({row["record_key"] for row in source_rows}) != len(source_rows):
        errors.append("record_key values are not unique")

    source_workbook = load_workbook(XLSX_SOURCE, data_only=True)
    opportunity_sheet = source_workbook["Opportunities"]
    opportunity_headers = {
        str(opportunity_sheet.cell(1, column).value): column
        for column in range(1, opportunity_sheet.max_column + 1)
        if opportunity_sheet.cell(1, column).value is not None
    }
    source_families = Counter()
    for row_number in range(2, opportunity_sheet.max_row + 1):
        family = opportunity_sheet.cell(row_number, opportunity_headers["source_family"]).value
        if family:
            source_families[str(family)] += 1
    registry = source_workbook["Source_Registry"]
    registry_headers = {
        str(registry.cell(1, column).value): column
        for column in range(1, registry.max_column + 1)
        if registry.cell(1, column).value is not None
    }
    registered_families = set()
    for row_number in range(2, registry.max_row + 1):
        for label in ("provider", "canonical_domain"):
            value = registry.cell(row_number, registry_headers[label]).value
            if value:
                registered_families.add(str(value).lower().removeprefix("www."))
    new_source_families = [
        family for family in source_families
        if family.lower().removeprefix("www.") not in registered_families
    ]

    rows: list[dict[str, str]] = []
    for original in source_rows:
        row = dict(original)
        row.update(
            {
                "agent_verdict": verdict_for(row),
                "agent_checked_at_utc": CHECKED_AT.isoformat().replace("+00:00", "Z"),
                "official_listing_url": "",
                "value_score_0_100": "",
                "effort_score_0_100": "",
                "risk_score_0_100": "",
                "evidence_files": "",
            }
        )
        rows.append(row)

    if not errors:
        with CSV_REVIEWED.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=headers + AGENT_COLUMNS, extrasaction="ignore", lineterminator="\n")
            writer.writeheader()
            writer.writerows(rows)

        workbook = load_workbook(XLSX_SOURCE)
        sheet = workbook["Opportunities"]
        existing_headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
        index_by_key = {row["record_key"]: row for row in rows}
        if len(existing_headers) < 43 or existing_headers[0] != "record_key":
            errors.append("Unexpected Opportunities worksheet headers")
        else:
            start_column = len(existing_headers) + 1
            for offset, label in enumerate(AGENT_COLUMNS):
                cell = sheet.cell(1, start_column + offset, label)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                sheet.column_dimensions[get_column_letter(start_column + offset)].width = max(16, len(label) + 2)
            row_headers = {str(sheet.cell(1, column).value): column for column in range(1, sheet.max_column + 1)}
            key_column = row_headers["record_key"]
            for row_number in range(2, sheet.max_row + 1):
                key = str(sheet.cell(row_number, key_column).value or "")
                if key not in index_by_key:
                    errors.append(f"Workbook row {row_number}: unknown record_key {key!r}")
                    continue
                result = index_by_key[key]
                for offset, label in enumerate(AGENT_COLUMNS):
                    sheet.cell(row_number, start_column + offset, result[label] or None)
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
            workbook.save(XLSX_REVIEWED)

    if errors:
        raise ValueError("Validation failed:\n" + "\n".join(errors[:50]))

    verdict_counts = Counter(row["agent_verdict"] for row in rows)
    action_counts = Counter(row["action"] for row in rows)
    status_counts = Counter(row["verification_status"] for row in rows)
    due_rows: list[tuple[datetime, str, str]] = []
    date_only_rows: list[tuple[date, str, str]] = []
    for row in rows:
        if row["agent_verdict"] != "unresolved":
            continue
        explicit = parse_time(row["deadline_utc"])
        day = parse_date(row["deadline_date"])
        if explicit and CHECKED_AT < explicit <= CHECKED_AT + timedelta(days=7):
            due_rows.append((explicit, row["record_key"], row["title"]))
        elif not explicit and day and CHECKED_AT.date() <= day <= (CHECKED_AT + timedelta(days=7)).date():
            date_only_rows.append((day, row["record_key"], row["title"]))

    data = {
        "generated_at_utc": CHECKED_AT.isoformat().replace("+00:00", "Z"),
        "as_of_date": CHECKED_AT.date().isoformat(),
        "rows_total": len(rows),
        "columns_source_interchange": len(headers),
        "columns_reviewed_csv": len(headers) + len(AGENT_COLUMNS),
        "unique_record_keys": len({row["record_key"] for row in rows}),
        "json_cells_parsed_ok": parsed_json_cells,
        "problems": errors,
        "actions": dict(action_counts),
        "verification_status": dict(status_counts),
        "agent_verdict": dict(verdict_counts),
        "import_ready_true": sum(row["import_ready"].lower() == "true" for row in rows),
        "verified_rows": sum(row["verification_status"] == "verified" for row in rows),
        "deadline_expired_as_of_delivery": sum(row["agent_verdict"] == "expired" for row in rows),
        "deadline_date_only_within_72h": sum(
            CHECKED_AT.date() <= (parse_date(row["deadline_date"]) or date.min) <= (CHECKED_AT + timedelta(days=3)).date()
            and not row["deadline_utc"] and row["agent_verdict"] == "unresolved"
            for row in rows
        ),
        "explicit_deadlines_within_72h": [
            {"record_key": key, "title": title, "deadline_utc": when.isoformat().replace("+00:00", "Z")}
            for when, key, title in due_rows if when <= CHECKED_AT + timedelta(days=3)
        ],
        "date_only_deadlines_within_7_days": [
            {"record_key": key, "title": title, "deadline_date": day.isoformat()}
            for day, key, title in date_only_rows
        ],
        "source_family_counts": dict(source_families),
        "new_source_families": new_source_families,
        "dedup_flags": json.loads((ROOT / "validation_report.json").read_text(encoding="utf-8")).get("dedup_flags", []),
        "sha256": {
            path.name: sha256(path)
            for path in (CSV_SOURCE, XLSX_SOURCE, README, CSV_REVIEWED, XLSX_REVIEWED)
        },
        "limitations": [
            "No provider-owned pages were rechecked; no row is promoted to verified.",
            "No database or live app writes were made, per the source README.",
            "All source rows have import_ready=false. Opportunity.organizationId is required by the current app schema, so these unlinked candidates remain in staging.",
            "Value, effort, and risk scores are blank because no scoring evidence or rubric was completed.",
            "Date-only deadlines are reported as calendar dates; no time or timezone was inferred.",
        ],
    }
    REPORT.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    explicit_72 = [entry for entry in due_rows if entry[0] <= CHECKED_AT + timedelta(days=3)]
    lines = [
        "# Opportunity intake handoff",
        "",
        f"Audit cutoff: {CHECKED_AT.isoformat().replace('+00:00', 'Z')} (12:15:53 EAT).",
        "",
        "## Outcome",
        "",
        f"- {len(rows)} rows reviewed offline: {action_counts['create']} creates, {action_counts['update_existing']} update, {action_counts['exclude']} exclusions.",
        f"- Verdicts: {verdict_counts['unresolved']} unresolved, {verdict_counts['expired']} expired, {verdict_counts['rejected']} rejected, {verdict_counts['verified']} verified.",
        f"- {data['import_ready_true']} rows marked import-ready; {data['verified_rows']} verified rows.",
        "- No live database/app writes were made. The current Opportunity model requires organizationId, so unattached records cannot be inserted into it. The ZIP README also prohibits live writes for this run.",
        "- The ZIP contained the workbook, 43-column CSV, README, and original validation report. All are extracted alongside this handoff.",
        "",
        "## Explicit deadlines within 72 hours",
        "",
    ]
    lines += [f"- {when.isoformat().replace('+00:00', 'Z')}: {title} [{key}]" for when, key, title in explicit_72] or ["- None recorded with a precise UTC deadline."]
    lines += ["", "## Deadline dates within 72 hours (time or timezone not stated)", ""]
    date_only_72h = [entry for entry in date_only_rows if entry[0] <= (CHECKED_AT + timedelta(days=3)).date()]
    lines += [f"- {day.isoformat()}: {title} [{key}]" for day, key, title in date_only_72h] or ["- None."]
    lines += ["", "## Deadline dates within 7 days (time or timezone not stated)", ""]
    lines += [f"- {day.isoformat()}: {title} [{key}]" for day, key, title in date_only_rows] or ["- None."]
    lines += ["", "## Source families in this run", ""]
    lines += [f"- {family}: {count}" for family, count in source_families.most_common()]
    lines += [f"- New relative to Source_Registry: {', '.join(new_source_families) if new_source_families else 'none'}"]
    lines += [
        "",
        "## Next action",
        "",
        "Review the short-window rows first on provider-owned pages. Keep unresolved records pending, preserve the dedup flags, and only map records into the app after an approved data-model path for organization-independent opportunities exists.",
        "",
        "See `validation_report.reviewed.json` for parse/enum/deadline checks and SHA-256 hashes. Scores and official listing URLs remain blank because the required evidence was not obtained.",
    ]
    HANDOFF.write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
